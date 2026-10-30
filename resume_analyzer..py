# AI Resume Analyzer & ATS Optimizer — V12.2
# Run: python -m streamlit run resume_analyzer.py
# This version avoids pandas and st.bar_chart to prevent the DLL issue.

import io
import re
from collections import Counter
import streamlit as st
from pypdf import PdfReader
from docx import Document
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

APP_VERSION = "V12.2"
st.set_page_config(page_title="ResumeAI — Resume Analyzer", page_icon="📄", layout="wide")

if st.session_state.get("app_version") != APP_VERSION:
    st.session_state["app_version"] = APP_VERSION
    st.session_state["analysis"] = None
if "analysis" not in st.session_state:
    st.session_state["analysis"] = None

st.markdown("""
<style>
.block-container{padding-top:1.5rem;max-width:1250px}
.hero{padding:1.4rem;border-radius:18px;background:linear-gradient(120deg,rgba(68,97,242,.15),rgba(24,180,160,.10));border:1px solid rgba(120,130,150,.25);margin-bottom:1rem}
.muted{color:#778096;font-size:.92rem}
.pill{display:inline-block;padding:.25rem .6rem;border-radius:999px;background:rgba(100,116,139,.14);margin:.15rem;font-size:.88rem}
</style>
""", unsafe_allow_html=True)

def norm(text):
    text = (text or "").lower().replace("c++", " cpp ").replace("c#", " csharp ")
    text = text.replace("node.js", "nodejs").replace("restful api", "rest api")
    text = re.sub(r"[\u2010-\u2015]", "-", text)
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9+#./\- ]+", " ", text)).strip()

def has_phrase(text, phrase):
    t, p = norm(text), norm(phrase)
    if not p: return False
    pattern = r"(?<![a-z0-9])" + re.escape(p).replace(r"\ ", r"\s+") + r"(?![a-z0-9])"
    return re.search(pattern, t) is not None

SKILL_GROUPS = {
    "Programming": ["Python","Java","JavaScript","C","C++","C#","PHP","Ruby","Go","Kotlin","Swift","TypeScript"],
    "Web Development": ["HTML","CSS","React","Node.js","REST API","Django","Flask","Spring Boot","Express.js","Angular","Vue.js"],
    "Database": ["SQL","MySQL","PostgreSQL","MongoDB","Firebase","Oracle","DBMS"],
    "Tools & Cloud": ["Git","GitHub","AWS","Azure","Google Cloud","Docker","Kubernetes","Linux","Jenkins"],
    "Software Engineering": ["Data Structures","Algorithms","OOP","Object-Oriented Programming","Problem Solving","System Design","Debugging","Unit Testing","Design Patterns","Operating Systems","Computer Networks"],
    "AI & Data": ["Artificial Intelligence","AI","Machine Learning","ML","Data Analysis","Data Science","Pandas","NumPy","TensorFlow","PyTorch"],
    "Design & Low-code": ["Mendix","UI/UX Design","Figma","Android Development","Android Studio"],
}
REQUIRED_DEFAULT = ["Python","Java","JavaScript","C++","OOP","Data Structures","Algorithms","Problem Solving","SQL","MySQL","REST API","Git","GitHub","Debugging","Unit Testing","DBMS","Operating Systems","Computer Networks"]
PREFERRED_DEFAULT = ["React","Django","Flask","Spring Boot","AWS","Docker","System Design","Machine Learning","Pandas","NumPy"]
STOPWORDS = set("""
a an the and or to of in for with on at by from as is are be will can should we our you your
they their this that these those it its role job candidate ideal about years year entry level
responsibilities responsibility preferred required skills skill ability abilities strong good
excellent related field degree equivalent practical including such other new learn learning
team teams cross functional motivated reliable clean efficient maintain maintainable develop
development design software computer science application applications work working build building
use using ensure following fundamentals principles contribute participate deliver solutions
technology technologies project projects code coding programming data systems system basic
knowledge understanding demonstrate demonstrates analytical communication effective relevant
resume resumes description teamwork ownership backend
""".split())

def extract_text(upload):
    if upload.name.lower().endswith(".pdf"):
        reader = PdfReader(upload)
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if upload.name.lower().endswith(".docx"):
        doc = Document(upload)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts).strip()
    raise ValueError("Upload a PDF or DOCX file.")

def get_skills(text):
    found = set()
    for group in SKILL_GROUPS.values():
        for skill in group:
            if has_phrase(text, skill):
                if skill == "AI" and has_phrase(text, "Artificial Intelligence"): continue
                if skill == "ML" and has_phrase(text, "Machine Learning"): continue
                found.add(skill)
    if has_phrase(text, "Object-Oriented Programming"): found.add("OOP")
    if has_phrase(text, "Data Structures and Algorithms") or has_phrase(text, "DSA"):
        found.update(["Data Structures", "Algorithms"])
    if has_phrase(text, "RESTful API") or has_phrase(text, "REST APIs"): found.add("REST API")
    if has_phrase(text, "Spring"): found.add("Spring Boot")
    return sorted(found, key=str.lower)

def classify_skills(jd):
    required_markers = ("required","must have","must-have","essential","minimum qualifications")
    preferred_markers = ("preferred","nice to have","nice-to-have","bonus","desired")
    required_text, preferred_text, section = [], [], "general"
    for line in jd.splitlines():
        low = line.lower()
        if any(x in low for x in required_markers): section = "required"
        elif any(x in low for x in preferred_markers): section = "preferred"
        if section == "required": required_text.append(line)
        elif section == "preferred": preferred_text.append(line)
    all_skills = set(get_skills(jd))
    required = set(get_skills("\n".join(required_text)))
    preferred = set(get_skills("\n".join(preferred_text)))
    if not required and not preferred:
        required = all_skills.intersection(REQUIRED_DEFAULT)
        preferred = all_skills.intersection(PREFERRED_DEFAULT)
    else:
        required -= preferred
    if not required: required = all_skills.intersection(REQUIRED_DEFAULT)
    if not preferred: preferred = all_skills.intersection(PREFERRED_DEFAULT) - required
    return sorted(required,key=str.lower), sorted(preferred,key=str.lower), sorted(all_skills,key=str.lower)

PHRASES = ["REST API","Data Structures and Algorithms","Object Oriented Programming","Problem Solving",
"System Design","Unit Testing","Design Patterns","Operating Systems","Computer Networks",
"Relational Databases","Version Control","Code Review","Backend Services","Cloud Computing",
"Application Performance","Database Management Systems","Spring Boot","Machine Learning",
"Artificial Intelligence","Debugging"]
def get_keywords(jd, limit=24):
    keywords = [p.title() for p in PHRASES if has_phrase(jd,p)]
    tokens = re.findall(r"\b[a-z][a-z0-9+#.-]{2,}\b", norm(jd))
    counts = Counter(t for t in tokens if t not in STOPWORDS and not t.isdigit())
    for token,_ in counts.most_common(limit*2):
        if token.lower() not in {x.lower() for x in keywords}:
            keywords.append(token)
        if len(keywords) >= limit: break
    result, seen = [], set()
    for k in keywords:
        if k.lower() not in seen:
            result.append(k); seen.add(k.lower())
    return result[:limit]

def get_sections(text):
    groups = {
        "Summary / Objective":["summary","professional summary","career objective","objective","profile"],
        "Skills":["technical skills","skills","core competencies"],
        "Experience":["experience","work experience","professional experience","employment","internship"],
        "Projects":["projects","academic projects","personal projects"],
        "Education":["education","academic background","qualification"],
        "Certifications":["certifications","certificates"],
        "Achievements":["achievements","awards","accomplishments"],
    }
    lines = [re.sub(r"[^a-z /&]","",x.strip().lower()) for x in text.splitlines()]
    return {name:any(line == v or line.startswith(v+" ") or (len(line)<45 and v in line) for line in lines for v in variants) for name,variants in groups.items()}

ACTION_VERBS = ["built","developed","designed","implemented","created","automated","optimized","improved","reduced","increased","delivered","deployed","integrated","analyzed","managed","led","collaborated","tested","maintained","engineered","configured","streamlined","resolved","launched","migrated","documented"]

def html_escape(value):
    return str(value or "").replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")

def content_stats(text):
    words = re.findall(r"\b[\w+#.-]+\b",text)
    low = text.lower()
    actions = sum(bool(re.search(r"\b"+v+r"\w*\b",low)) for v in ACTION_VERBS)
    metrics = re.findall(r"\b\d+(?:\.\d+)?\s?(?:%|percent|users|customers|records|hours|days|weeks|months|years|projects|teams|ms|seconds|minutes)\b",low)
    return len(words),actions,len(metrics)

def make_summary(skills, matches, jd):
    role = "software development"
    m = re.search(r"(software development engineer|software engineer|java developer|python developer|backend developer|full.stack developer|mendix developer)",jd,re.I)
    if m: role = m.group(0)
    useful = [x for x in ["Java","Python","JavaScript","SQL","MySQL","React","Mendix","AWS"] if x in skills]
    if not useful: useful = sorted(skills,key=str.lower)[:5]
    return f"Software development candidate targeting {role} opportunities, with exposure to {', '.join(useful) if useful else 'programming fundamentals and software projects'}. Brings practical project experience and a focus on maintainable code, problem-solving, and reliable applications. Relevant matched skills include {', '.join(matches[:5] or useful) if (matches or useful) else 'relevant technical skills'}."

def build_projects(jdskills):
    suggestions=[]
    if set(jdskills)&{"REST API","Spring Boot","Django","Flask"}:
        suggestions.append("Build a documented REST API with validation, authentication, database integration, tests, and API documentation.")
    if set(jdskills)&{"AWS","Docker","Kubernetes"}:
        suggestions.append("Deploy an existing project with a simple CI workflow and a README explaining architecture and setup.")
    if set(jdskills)&{"SQL","MySQL","PostgreSQL"}:
        suggestions.append("Add thoughtful schema design, joins, error handling, and database tests to an existing project.")
    if set(jdskills)&{"Data Structures","Algorithms","Problem Solving"}:
        suggestions.append("Create a small DSA repository with tested solutions and short complexity explanations.")
    return suggestions or ["Choose a project aligned to the target job and document the design, setup, tests, and your individual contribution."]

def technical_questions(skills):
    bank={
    "Java":["Explain interfaces vs abstract classes in Java.","How do HashMap and HashSet work?"],
    "Python":["Explain mutable vs immutable objects in Python.","How do you handle exceptions in a maintainable Python app?"],
    "SQL":["Explain INNER JOIN vs LEFT JOIN.","How would you investigate a slow SQL query?"],
    "MySQL":["What are indexes and their trade-offs?","Explain transactions and ACID properties."],
    "REST API":["Explain HTTP methods and common status codes.","How would you validate and secure an API?"],
    "Spring Boot":["Explain dependency injection in Spring.","How would you test a Spring Boot endpoint?"],
    "Django":["Explain Django's model-view-template pattern.","How would you secure a Django endpoint?"],
    "Flask":["How would you structure a growing Flask app?","How do you validate requests and handle errors?"],
    "React":["Explain props vs state.","How do you handle API loading and error states?"],
    "Data Structures":["When would you use a hash map instead of a list?","How do you choose a data structure for fast lookups?"],
    "Algorithms":["Explain O(n) vs O(log n).","How would you find duplicates efficiently?"],
    "OOP":["Explain encapsulation, inheritance, polymorphism, and abstraction.","When is composition better than inheritance?"],
    "Problem Solving":["How do you break down an unfamiliar coding problem?","How do you test edge cases and explain complexity?"],
    "AWS":["How would you deploy a simple application to AWS?","How do you protect cloud credentials?"],
    "Git":["Explain merge vs rebase.","How would you resolve a merge conflict?"],
    "Mendix":["Explain domain models and microflows in Mendix.","How do you implement role-based access?"]}
    result=[]
    for key,qs in bank.items():
        if key in skills: result.extend(qs)
    return result[:12] or ["Explain a project you built and the trade-offs you made.","How do you test and debug code?","Describe a data structure or algorithm you used."]

def interview_plan(missing):
    focus=", ".join(missing[:3]) if missing else "core programming and project explanations"
    return [
    ("Day 1 — Resume","Prepare a 60–90 second introduction and explain every resume bullet honestly."),
    ("Day 2 — Coding",f"Practice arrays, strings, hash maps, complexity, and problems related to {focus}."),
    ("Day 3 — Core CS","Review OOP, SQL/DBMS, operating systems, and networking relevant to the role."),
    ("Day 4 — Projects","Prepare architecture, design decisions, edge cases, testing, and lessons from two projects."),
    ("Day 5 — Mock interview","Complete a timed coding session and practice behavioral answers with specific examples.")]

def pdf_report(data, interview=False):
    buffer=io.BytesIO()
    doc=SimpleDocTemplate(buffer,pagesize=A4,rightMargin=40,leftMargin=40,topMargin=42,bottomMargin=42)
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle2",parent=styles["Title"],alignment=TA_CENTER,textColor=colors.HexColor("#243b75")))
    styles.add(ParagraphStyle(name="Section2",parent=styles["Heading2"],textColor=colors.HexColor("#243b75"),spaceBefore=12,spaceAfter=6))
    story=[Paragraph("ResumeAI Interview Preparation Report" if interview else "ResumeAI ATS Analysis Report",styles["ReportTitle2"]),
           Paragraph("<b>Resume:</b> "+html_escape(data["filename"]),styles["BodyText"]),Spacer(1,8)]
    if not interview:
        rows=[["Metric","Score"],["Overall ATS",f"{data['overall']}%"],["Job Match",f"{data['job_match']}%"],["Skill Match",f"{data['skill_score']}%"],["Keyword Match",f"{data['keyword_score']}%"],["Structure",f"{data['structure_score']}%"],["Content",f"{data['content_score']}%"]]
        table=Table(rows,colWidths=[300,130])
        table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#243b75")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.5,colors.lightgrey),("PADDING",(0,0),(-1,-1),7)]))
        story.append(table)
        sections=[("Matching required skills",data["match_req"]),("Missing required skills",data["miss_req"]),("Matching keywords",data["match_kw"]),("Missing keywords",data["miss_kw"]),("Strengths",data["strengths"]),("Areas to improve",data["weaknesses"]),("Recommendations",data["recommendations"])]
    else:
        sections=[("Strengths",data["strengths"]),("Areas to improve",data["weaknesses"]),("Skills to prioritize",data["miss_req"] or data["miss_pref"] or ["Continue role-specific practice"]),("Technical questions",data["tech_questions"]),("HR questions",data["hr_questions"]),("Five-day preparation plan",[a+": "+b for a,b in data["plan"]])]
    for heading,items in sections:
        story.append(Paragraph(html_escape(heading),styles["Section2"]))
        for item in items or ["None detected"]:
            story.append(Paragraph("• "+html_escape(item),styles["BodyText"])); story.append(Spacer(1,3))
    story.append(Spacer(1,12))
    story.append(Paragraph("Disclaimer: Scores are heuristic estimates, not official employer ATS scores or hiring predictions. Verify all generated suggestions.",styles["Italic"]))
    doc.build(story); buffer.seek(0); return buffer.getvalue()

def pill_list(items, empty="None detected"):
    if items: st.markdown(" ".join(f'<span class="pill">{html_escape(x)}</span>' for x in items),unsafe_allow_html=True)
    else: st.caption(empty)

def bar(label,value):
    st.write(f"**{label}: {value}%**"); st.progress(max(0,min(100,int(value))))

st.markdown('<div class="hero"><h1>📄 ResumeAI — Resume Analyzer & ATS Optimizer</h1><div class="muted">Analyze a resume against a job description, identify gaps, and prepare for interviews.</div></div>',unsafe_allow_html=True)
with st.sidebar:
    st.markdown("### ⚙️ Analysis")
    st.caption("Version "+APP_VERSION)
    st.markdown("1. Upload PDF/DOCX.\n2. Paste job description.\n3. Click Analyze.\n4. Review results and reports.")
    st.info("Use resumes only with the person's permission. This app does not intentionally save uploaded files to a database.")
    if st.button("🧹 Reset analysis",use_container_width=True):
        st.session_state["analysis"]=None; st.rerun()

c1,c2=st.columns(2)
with c1: upload=st.file_uploader("📎 Upload resume (PDF or DOCX)",type=["pdf","docx"])
with c2: jd=st.text_area("🎯 Paste job description",height=170,placeholder="Include responsibilities, required skills, and preferred skills...")
if st.button("🔍 Analyze Resume",type="primary",use_container_width=True):
    if not upload:
        st.error("Please upload a resume first.")
    elif not jd.strip():
        st.error("Please paste a job description first.")
    else:
        try:
            resume=extract_text(upload)
            if len(resume.strip())<30:
                st.error("Very little text could be extracted. If this is a scanned PDF, try a text-based PDF or DOCX.")
            else:
                skills=set(get_skills(resume))
                req,pref,jskills=classify_skills(jd)
                match_req=sorted(skills&set(req),key=str.lower)
                miss_req=sorted(set(req)-skills,key=str.lower)
                match_pref=sorted(skills&set(pref),key=str.lower)
                miss_pref=sorted(set(pref)-skills,key=str.lower)
                weight=len(req)*3+len(pref)
                skill_score=round(100*(len(match_req)*3+len(match_pref))/weight) if weight else 100
                kws=get_keywords(jd)
                match_kw=[k for k in kws if has_phrase(resume,k)]
                miss_kw=[k for k in kws if k not in match_kw]
                keyword_score=round(100*len(match_kw)/len(kws)) if kws else 100
                sections=get_sections(resume)
                structure_score=round(100*sum(sections.get(x,False) for x in ["Summary / Objective","Skills","Experience","Projects","Education"])/5)
                word_count,action_count,metrics_count=content_stats(resume)
                content_score=min(100,45+(15 if word_count>=250 else 10 if word_count>=120 else 0)+(10 if sections.get("Experience") else 0)+(10 if sections.get("Projects") else 0)+(10 if action_count>=5 else 5 if action_count>=2 else 0)+(10 if metrics_count else 0))
                overall=round(skill_score*.40+keyword_score*.25+structure_score*.20+content_score*.15)
                job_match=round(skill_score*.70+keyword_score*.30)
                strengths=[]
                if match_req: strengths.append("Required skills detected: "+", ".join(match_req[:8])+".")
                if sections.get("Projects"): strengths.append("A projects section was detected.")
                if sections.get("Experience"): strengths.append("An experience section was detected.")
                if sections.get("Education"): strengths.append("An education section was detected.")
                if metrics_count: strengths.append(f"{metrics_count} measurable result(s) detected.")
                if not strengths: strengths.append("Resume text extracted successfully; review the detailed findings.")
                weaknesses=[]
                if miss_req: weaknesses.append("Required skills not detected: "+", ".join(miss_req)+". Add them only if you genuinely have them.")
                if keyword_score<60: weaknesses.append("Keyword match is low; use relevant job-description wording naturally.")
                if not metrics_count: weaknesses.append("No measurable outcomes were detected; add genuine metrics where available.")
                if not sections.get("Experience"): weaknesses.append("A clearly labeled experience section was not detected.")
                if not sections.get("Projects"): weaknesses.append("A clearly labeled projects section was not detected.")
                if not weaknesses: weaknesses.append("No major heuristic issues detected; tailor this resume to the exact vacancy.")
                recommendations=[]
                if miss_req: recommendations.append("Review missing required skills and add only skills you can confidently explain.")
                recommendations += ["Use action verbs and explain the action, technology, and result in each bullet.","Add truthful measurable outcomes where possible.","Use relevant job-description terms only when they accurately describe your experience.","Keep standard section headings and readable bullets."]
                summary=make_summary(skills,match_req,jd)
                original_lines=[x.strip() for x in resume.splitlines() if x.strip()]
                candidate_name=original_lines[0] if original_lines else "Candidate Name"
                contact=original_lines[1] if len(original_lines)>1 and ("@" in original_lines[1] or re.search(r"\+?\d[\d ()-]{7,}",original_lines[1])) else ""
                role_match=re.search(r"(software development engineer|software engineer|java developer|python developer|backend developer|full.stack developer|mendix developer)",jd,re.I)
                role=role_match.group(0).title() if role_match else "Target Role"
                optimized=f"{candidate_name}\n{contact}\nTARGET ROLE: {role}\n\nPROFESSIONAL SUMMARY\n{summary}\n\nTECHNICAL SKILLS\n{', '.join(sorted(skills,key=str.lower))}\n\nORIGINAL RESUME CONTENT\n{resume}\n\nDraft only: verify details, remove duplicates, and never add claims you cannot substantiate."
                st.session_state["analysis"]={
                    "filename":upload.name,"resume":resume,"jd":jd,"overall":overall,"job_match":job_match,
                    "skill_score":skill_score,"keyword_score":keyword_score,"structure_score":structure_score,"content_score":content_score,
                    "skills":sorted(skills,key=str.lower),"job_skills":jskills,"req":req,"pref":pref,
                    "match_req":match_req,"miss_req":miss_req,"match_pref":match_pref,"miss_pref":miss_pref,
                    "match_kw":match_kw,"miss_kw":miss_kw,"sections":sections,"word_count":word_count,
                    "action_count":action_count,"metrics_count":metrics_count,"strengths":strengths,"weaknesses":weaknesses,
                    "recommendations":recommendations,"summary":summary,"optimized":optimized,
                    "projects":build_projects(jskills),"tech_questions":technical_questions(set(jskills)|set(miss_req)),
                    "hr_questions":["Tell me about yourself.","Why are you interested in this role?","Describe a technical challenge you solved.","Tell me about teamwork on a project.","Describe a mistake and what you learned.","How do you learn a new technology?","Why should we hire you?","What are your career goals?"],
                    "plan":interview_plan(miss_req)
                }
                st.success("Analysis complete!")
        except Exception as exc:
            st.error(f"Could not analyze this file: {exc}")

data=st.session_state.get("analysis")
if data:
    st.markdown("---")
    st.subheader("📊 Resume Analytics")
    a,b,c,d=st.columns(4)
    a.metric("ATS Score",f"{data['overall']}%"); b.metric("Job Match",f"{data['job_match']}%")
    c.metric("Skill Match",f"{data['skill_score']}%"); d.metric("Keyword Match",f"{data['keyword_score']}%")
    st.caption("These are heuristic estimates, not official ATS scores or hiring predictions.")
    tabs=st.tabs(["📊 Overview","🎯 Job Match","🧠 Skills","✍️ AI Improvement","🚀 Projects","📄 Optimized Resume","🎤 Interview Prep","📋 Structure","📄 Resume Text"])
    with tabs[0]:
        for label,key in [("Skill Match","skill_score"),("Keyword Match","keyword_score"),("Structure","structure_score"),("Content","content_score")]: bar(label,data[key])
        x,y,z=st.columns(3); x.metric("Resume words",data["word_count"]); y.metric("Action verbs detected",data["action_count"]); z.metric("Metrics detected",data["metrics_count"])
        st.subheader("Strengths")
        for item in data["strengths"]: st.markdown("- ✅ "+item)
        st.subheader("Areas to improve")
        for item in data["weaknesses"]: st.markdown("- ⚠️ "+item)
        st.download_button("⬇️ Download Professional PDF Report",data=pdf_report(data),file_name="ResumeAI_ATS_Analysis_Report.pdf",mime="application/pdf",use_container_width=True)
    with tabs[1]:
        st.subheader(f"Job Match Score: {data['job_match']}%"); st.progress(data["job_match"])
        st.markdown("**Matching required skills**"); pill_list(data["match_req"])
        st.markdown("**Missing required skills**"); pill_list(data["miss_req"],"No major required skill gaps detected.")
        st.markdown("**Matching preferred skills**"); pill_list(data["match_pref"])
        st.markdown("**Missing preferred skills**"); pill_list(data["miss_pref"])
        st.caption("A skill may be marked missing if its wording was not detected; this does not prove the candidate lacks it.")
    with tabs[2]:
        for group,group_skills in SKILL_GROUPS.items():
            found=[x for x in group_skills if x in data["skills"]]
            if found: st.markdown("**"+group+"**"); pill_list(found)
        st.subheader("Skills to prioritize"); pill_list(data["miss_req"],"No major required skill gaps detected.")
    with tabs[3]:
        st.subheader("Recommendations")
        for item in data["recommendations"]: st.markdown("- 💡 "+item)
        st.subheader("Suggested professional summary"); st.write(data["summary"])
        st.caption("Edit this summary to ensure it truthfully reflects the candidate.")
        bullet=st.text_area("Paste one existing resume bullet",placeholder="Example: Worked on an application using Java and MySQL.",key="bullet_input")
        if st.button("✨ Improve this bullet"):
            if not bullet.strip(): st.warning("Paste a bullet first.")
            else: st.info("Use this structure: **Action verb + task + technology + verified result.**\n\nDraft to refine: "+bullet.strip()+"\n\nDo not invent metrics or achievements.")
    with tabs[4]:
        st.subheader("Project recommendations")
        for item in data["projects"]: st.markdown("- 🚀 "+item)
        st.markdown("**Project explanation checklist**\n- Problem solved\n- Your individual contribution\n- Technology choices\n- Hardest bug or trade-off\n- Testing approach\n- Future improvements")
    with tabs[5]:
        st.warning("Draft only. Verify all details, remove duplicates, and never submit claims you cannot substantiate.")
        edited=st.text_area("Editable optimized resume draft",value=data["optimized"],height=420,key="optimized_editor")
        st.download_button("⬇️ Download Optimized Resume Draft (.txt)",data=edited,file_name="ResumeAI_Optimized_Resume_Draft.txt",mime="text/plain",use_container_width=True)
    with tabs[6]:
        st.subheader("Technical interview questions")
        for i,q in enumerate(data["tech_questions"],1): st.markdown(f"**{i}.** {q}")
        st.subheader("HR / behavioral questions")
        for i,q in enumerate(data["hr_questions"],1): st.markdown(f"**{i}.** {q}")
        st.subheader("Five-day preparation plan")
        for day,task in data["plan"]: st.markdown(f"**{day}** — {task}")
        st.download_button("⬇️ Download Interview Preparation Report",data=pdf_report(data,interview=True),file_name="ResumeAI_Interview_Preparation_Report.pdf",mime="application/pdf",use_container_width=True)
    with tabs[7]:
        for section,present in data["sections"].items(): st.markdown(f"{'✅' if present else '⚠️'} **{section}** — {'Detected' if present else 'Not detected'}")
        bar("Structure score",data["structure_score"])
    with tabs[8]:
        st.text_area("Extracted resume text",value=data["resume"],height=420)
        st.download_button("⬇️ Download Extracted Text",data=data["resume"],file_name="Extracted_Resume_Text.txt",mime="text/plain",use_container_width=True)
else:
    st.markdown("---")
    st.subheader("✨ What this app can do")
    a,b,c=st.columns(3)
    a.info("🎯 Job Match\n\nCompare resume skills with the job description.")
    b.info("🧠 Resume Insights\n\nReview keywords, structure, strengths, and gaps.")
    c.info("🎤 Interview Prep\n\nPractice technical and HR questions.")
    st.info("Upload a resume and paste a job description above to begin.")

