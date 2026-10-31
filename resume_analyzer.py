import streamlit as st
import re
import io

from pypdf import PdfReader
from docx import Document

from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER


# ============================================================
# APP VERSION
# ============================================================

APP_VERSION = "V12.1"


# ============================================================
# SESSION STATE
# ============================================================

if st.session_state.get("app_version") != APP_VERSION:
    st.session_state["app_version"] = APP_VERSION
    st.session_state["analysis"] = None

if "analysis" not in st.session_state:
    st.session_state["analysis"] = None


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="ResumeAI - ATS Analyzer",
    page_icon="🚀",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ============================================================
# CSS
# ============================================================

st.markdown(
    """
    <style>

    .main {
        background-color: #f7f9fc;
    }

    .block-container {
        padding-top: 1.5rem;
        padding-bottom: 3rem;
    }

    .hero {
        padding: 28px;
        border-radius: 18px;
        background: linear-gradient(135deg, #111827, #1f2937);
        color: white;
        margin-bottom: 25px;
    }

    .hero h1 {
        font-size: 38px;
        margin-bottom: 5px;
    }

    .hero p {
        color: #d1d5db;
        font-size: 17px;
    }

    .success-box {
        padding: 18px;
        border-radius: 14px;
        background: #ecfdf5;
        border: 1px solid #a7f3d0;
        margin-bottom: 12px;
    }

    .warning-box {
        padding: 18px;
        border-radius: 14px;
        background: #fffbeb;
        border: 1px solid #fde68a;
        margin-bottom: 12px;
    }

    .danger-box {
        padding: 18px;
        border-radius: 14px;
        background: #fef2f2;
        border: 1px solid #fecaca;
        margin-bottom: 12px;
    }

    .info-box {
        padding: 18px;
        border-radius: 14px;
        background: #eff6ff;
        border: 1px solid #bfdbfe;
        margin-bottom: 12px;
    }

    .skill-box {
        padding: 14px 18px;
        border-radius: 12px;
        background: white;
        border: 1px solid #e5e7eb;
        margin-bottom: 10px;
    }

    .question-card {
        background: white;
        border: 1px solid #e5e7eb;
        border-radius: 14px;
        padding: 18px;
        margin-bottom: 12px;
    }

    .section-title {
        font-size: 25px;
        font-weight: 700;
        margin-top: 15px;
        margin-bottom: 15px;
        color: #111827;
    }

    .small-muted {
        color: #6b7280;
        font-size: 13px;
    }

    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# HEADER
# ============================================================

st.markdown(
    """
    <div class="hero">
        <h1>🚀 ResumeAI</h1>
        <p>AI Resume Analyzer • ATS Optimizer • Job Match • Interview Intelligence</p>
    </div>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# SKILLS
# ============================================================

SKILLS = [
    "python", "java", "c", "c++", "c#", "javascript",
    "typescript", "html", "css", "react", "angular", "vue",
    "node.js", "node", "express", "django", "flask", "fastapi",
    "spring", "sql", "mysql", "postgresql", "mongodb", "oracle",
    "sqlite", "git", "github", "gitlab", "docker", "kubernetes",
    "aws", "azure", "gcp", "linux", "rest api", "api", "json",
    "xml", "oop", "object-oriented programming", "data structures",
    "algorithms", "dsa", "problem solving", "debugging", "testing",
    "unit testing", "machine learning", "ml",
    "artificial intelligence", "ai", "deep learning",
    "data science", "data analysis", "pandas", "numpy",
    "matplotlib", "tensorflow", "pytorch", "scikit-learn",
    "nlp", "power bi", "tableau", "excel",
    "communication", "leadership", "teamwork", "time management"
]


CATEGORIES = {
    "Programming": [
        "python", "java", "c", "c++", "c#", "javascript", "typescript"
    ],
    "Web Development": [
        "html", "css", "react", "angular", "vue", "node.js",
        "node", "express", "django", "flask", "fastapi", "spring"
    ],
    "Database": [
        "sql", "mysql", "postgresql", "mongodb", "oracle", "sqlite"
    ],
    "Tools & Cloud": [
        "git", "github", "gitlab", "docker", "kubernetes",
        "aws", "azure", "gcp", "linux"
    ],
    "Software Engineering": [
        "rest api", "api", "json", "xml", "oop",
        "object-oriented programming", "data structures",
        "algorithms", "dsa", "problem solving", "debugging",
        "testing", "unit testing"
    ],
    "AI & Data": [
        "machine learning", "ml", "artificial intelligence", "ai",
        "deep learning", "data science", "data analysis",
        "pandas", "numpy", "matplotlib", "tensorflow",
        "pytorch", "scikit-learn", "nlp", "power bi",
        "tableau", "excel"
    ],
    "Soft Skills": [
        "communication", "leadership", "teamwork", "time management"
    ],
}


ALIASES = {
    "dsa": "data structures",
    "algorithm": "algorithms",
    "oop": "oop",
    "object oriented programming": "oop",
    "object-oriented programming": "oop",
    "artificial intelligence": "ai",
    "machine learning": "ml",
    "node": "node.js",
    "restful api": "rest api",
    "restful apis": "rest api",
    "rest apis": "rest api",
    "problem-solving": "problem solving",
    "problem solving skills": "problem solving",
}


STOP_WORDS = {
    "the", "and", "for", "with", "from", "that", "this", "are",
    "you", "your", "our", "will", "have", "has", "had", "job",
    "role", "work", "working", "candidate", "candidates",
    "developer", "development", "experience", "years", "year",
    "skills", "skill", "strong", "good", "knowledge", "ability",
    "required", "preferred", "responsibilities", "responsibility",
    "including", "etc", "using", "used", "use", "must", "should",
    "would", "company", "team", "position", "looking", "seeking",
    "join", "build", "building", "application", "applications"
}


GENERIC_WORDS = {
    "software", "engineering", "engineer", "technology",
    "technologies", "business", "professional", "solutions",
    "system", "systems", "project", "projects", "development",
    "developer", "web", "programming", "technical", "tools",
    "environment", "platform", "service", "services", "product",
    "products"
}


REQUIRED_WORDS = {
    "required", "must", "mandatory", "essential",
    "need", "needs", "should have", "required skills",
    "requirements"
}


PREFERRED_WORDS = {
    "preferred", "prefer", "nice to have", "bonus",
    "plus", "desirable", "good to have", "additional"
}


# ============================================================
# INTERVIEW QUESTIONS
# ============================================================

TECHNICAL_QUESTIONS = {

    "python": [
        "What are the main features of Python?",
        "What is the difference between a list, tuple, set and dictionary?",
        "Explain mutable and immutable objects in Python.",
        "What is a Python decorator?",
        "What is the difference between == and is?",
        "How does exception handling work in Python?"
    ],

    "java": [
        "What are the main principles of OOP in Java?",
        "What is the difference between JDK, JRE and JVM?",
        "Explain method overloading and overriding.",
        "What is inheritance in Java?",
        "What is the difference between an interface and abstract class?"
    ],

    "c++": [
        "What is the difference between a class and an object?",
        "Explain inheritance and polymorphism.",
        "What are pointers and references?",
        "What is the difference between stack and heap memory?",
        "What are constructors and destructors?"
    ],

    "sql": [
        "What is the difference between WHERE and HAVING?",
        "Explain INNER JOIN and LEFT JOIN.",
        "What is normalization?",
        "What is a primary key?",
        "What is the difference between DELETE, DROP and TRUNCATE?",
        "Write a query to find the second-highest salary."
    ],

    "mysql": [
        "What is MySQL?",
        "How do indexes improve database performance?",
        "What is a foreign key?",
        "Explain database normalization.",
        "What is the difference between CHAR and VARCHAR?"
    ],

    "git": [
        "What is Git and why is it used?",
        "What is the difference between git pull and git fetch?",
        "What is git commit?",
        "What is a Git branch?",
        "How would you resolve a merge conflict?"
    ],

    "github": [
        "What is GitHub and how is it different from Git?",
        "How do pull requests work?",
        "What is a repository?",
        "How would you collaborate using GitHub?"
    ],

    "rest api": [
        "What is a REST API?",
        "What HTTP methods are commonly used?",
        "What is the difference between GET and POST?",
        "What is an HTTP status code?",
        "What does stateless mean in REST?"
    ],

    "django": [
        "What is Django?",
        "Explain Django's MVT architecture.",
        "What is a Django model?",
        "What is Django ORM?",
        "How does URL routing work?"
    ],

    "flask": [
        "What is Flask?",
        "How is Flask different from Django?",
        "How do routes work in Flask?",
        "How would you create a REST API using Flask?"
    ],

    "react": [
        "What is React?",
        "What are components?",
        "What is the difference between props and state?",
        "What are React hooks?",
        "What is useEffect?"
    ],

    "javascript": [
        "What is the difference between var, let and const?",
        "Explain closures.",
        "What is the difference between == and ===?",
        "What are promises?",
        "What is asynchronous programming?"
    ],

    "html": [
        "What is semantic HTML?",
        "What is the difference between div and section?",
        "What are HTML forms?",
        "Why are meta tags important?"
    ],

    "css": [
        "What is the CSS box model?",
        "What is the difference between Flexbox and Grid?",
        "Explain CSS specificity.",
        "What are media queries?"
    ],

    "oop": [
        "What are the four pillars of OOP?",
        "Explain encapsulation.",
        "What is inheritance?",
        "What is polymorphism?",
        "What is abstraction?"
    ],

    "data structures": [
        "What is the difference between an array and linked list?",
        "Explain stacks and queues.",
        "What is a hash table?",
        "What is a binary search tree?",
        "When would you use a graph?"
    ],

    "algorithms": [
        "What is time complexity?",
        "Explain binary search.",
        "What is the difference between BFS and DFS?",
        "Explain common sorting algorithms.",
        "What is Big-O notation?"
    ],

    "problem solving": [
        "How do you approach a programming problem?",
        "How do you debug incorrect output?",
        "How do you optimize a slow solution?",
        "How do you choose a data structure?"
    ],

    "machine learning": [
        "What is machine learning?",
        "What is the difference between supervised and unsupervised learning?",
        "What is overfitting?",
        "What is train-test splitting?",
        "How do you evaluate a model?"
    ],

    "ml": [
        "What is machine learning?",
        "Explain supervised and unsupervised learning.",
        "What is overfitting?",
        "What is classification?",
        "What is regression?"
    ],

    "ai": [
        "What is artificial intelligence?",
        "What is the difference between AI and ML?",
        "What are real-world applications of AI?",
        "What is generative AI?"
    ],

    "pandas": [
        "What is Pandas?",
        "What is a DataFrame?",
        "How do you handle missing values?",
        "How do you filter rows in a DataFrame?"
    ],

    "numpy": [
        "What is NumPy?",
        "Why are NumPy arrays useful?",
        "What is vectorization?"
    ],

    "docker": [
        "What is Docker?",
        "What is a container?",
        "What is the difference between an image and container?",
        "Why is Docker useful?"
    ],

    "aws": [
        "What is AWS?",
        "What is EC2?",
        "What is S3?",
        "What is an availability zone?"
    ],

    "linux": [
        "What is Linux?",
        "What Linux commands have you used?",
        "What are Linux file permissions?"
    ]
}


HR_QUESTIONS = [
    "Tell me about yourself.",
    "Walk me through your resume.",
    "Why are you interested in this role?",
    "Why should we hire you?",
    "What are your greatest strengths?",
    "What is one weakness you are improving?",
    "Tell me about a challenging project.",
    "Tell me about a difficult problem you solved.",
    "How do you handle pressure?",
    "How do you work in a team?",
    "Where do you see yourself in five years?",
    "Why do you want to work for our company?",
    "What motivates you?",
    "Tell me about a failure and what you learned.",
    "Do you have any questions for us?"
]


# ============================================================
# FILE EXTRACTION
# ============================================================

def extract_pdf_text(file):
    try:
        reader = PdfReader(file)
        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )
    except Exception as e:
        return f"PDF extraction error: {e}"


def extract_docx_text(file):
    try:
        doc = Document(file)
        parts = []

        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text)

        for table in doc.tables:
            for row in table.rows:
                text = " ".join(
                    cell.text for cell in row.cells
                )
                if text.strip():
                    parts.append(text)

        return "\n".join(parts)

    except Exception as e:
        return f"DOCX extraction error: {e}"


def extract_text(file):
    name = file.name.lower()

    if name.endswith(".pdf"):
        return extract_pdf_text(file)

    if name.endswith(".docx"):
        return extract_docx_text(file)

    if name.endswith(".txt"):
        return file.read().decode(
            "utf-8",
            errors="ignore"
        )

    return ""


# ============================================================
# TEXT HELPERS
# ============================================================

def normalize_text(text):
    text = text.lower()

    # Replace whole words/phrases only, longest alias first, so
    # "node.js" never becomes "node.js.js" and "algorithms" never
    # becomes "algorithmss".
    for alias in sorted(ALIASES, key=len, reverse=True):
        canonical = ALIASES[alias]

        if alias == canonical:
            continue

        text = re.sub(
            rf"(?<![a-z0-9.]){re.escape(alias)}(?![a-z0-9]|\.js)",
            canonical,
            text
        )

    return re.sub(r"\s+", " ", text).strip()


def phrase_present(text, phrase):
    text = normalize_text(text)
    phrase = normalize_text(phrase)

    # "+" and "#" are blocked after a match so that the skill "c"
    # is not detected inside "c++" or "c#".
    return bool(
        re.search(
            rf"(?<![a-z0-9]){re.escape(phrase)}(?![a-z0-9+#])",
            text,
            re.IGNORECASE
        )
    )


# Skills whose display name is not simply Title Case.
DISPLAY_NAMES = {
    "oop": "OOP", "sql": "SQL", "mysql": "MySQL",
    "postgresql": "PostgreSQL", "mongodb": "MongoDB",
    "sqlite": "SQLite", "html": "HTML", "css": "CSS",
    "api": "API", "rest api": "REST API", "json": "JSON",
    "xml": "XML", "ai": "AI", "ml": "ML", "nlp": "NLP",
    "dsa": "DSA", "aws": "AWS", "gcp": "GCP", "c": "C",
    "c++": "C++", "c#": "C#", "node.js": "Node.js",
    "javascript": "JavaScript", "typescript": "TypeScript",
    "github": "GitHub", "gitlab": "GitLab",
    "fastapi": "FastAPI", "tensorflow": "TensorFlow",
    "pytorch": "PyTorch", "scikit-learn": "scikit-learn",
    "numpy": "NumPy", "power bi": "Power BI",
    "data structures": "Data Structures",
}


def display_skill(skill):
    skill = str(skill)
    return DISPLAY_NAMES.get(skill.lower(), skill.title())


def clean_for_keywords(text):
    text = normalize_text(text)
    text = re.sub(
        r"[^a-z0-9+#.\-\s]",
        " ",
        text
    )
    return re.sub(r"\s+", " ", text)


# ============================================================
# SKILLS
# ============================================================

def extract_skills(text):
    found = []
    normalized = normalize_text(text)

    for skill in SKILLS:
        if phrase_present(normalized, skill):
            canonical = ALIASES.get(skill, skill)

            if canonical not in found:
                found.append(canonical)

    return sorted(found)


def categorize_skills(skills):
    result = {}

    for category, category_skills in CATEGORIES.items():

        matches = []

        for skill in skills:
            canonical = ALIASES.get(skill, skill)

            if canonical in category_skills:
                matches.append(skill)

        if matches:
            result[category] = sorted(set(matches))

    return result


# ============================================================
# REQUIRED / PREFERRED
# ============================================================

def _last_hit(text, words):
    """Position of the closest keyword before the skill, or -1."""
    best = -1

    for word in words:
        position = text.rfind(word)
        best = max(best, position)

    return best


def detect_required_preferred_skills(jd_text, jd_skills):

    normalized = normalize_text(jd_text)

    required = []
    preferred = []

    for skill in jd_skills:

        matches = list(
            re.finditer(
                rf"(?<![a-z0-9]){re.escape(skill)}(?![a-z0-9+#])",
                normalized
            )
        )

        if not matches:
            continue

        kinds = []

        for match in matches:

            before = normalized[
                max(0, match.start() - 180):match.start()
            ]

            required_hit = _last_hit(before, REQUIRED_WORDS)
            preferred_hit = _last_hit(before, PREFERRED_WORDS)

            # The closest cue wins, so "Required: python ...
            # Preferred: aws" puts aws under preferred. A skill with
            # no cue is treated as required.
            if preferred_hit > required_hit:
                kinds.append("preferred")
            else:
                kinds.append("required")

        # One required mention is enough to make it required.
        if "required" in kinds:
            required.append(skill)
        else:
            preferred.append(skill)

    if not required:
        required = [
            skill for skill in jd_skills
            if skill not in preferred
        ]

    required = sorted(set(required))
    preferred = sorted(
        set(preferred) - set(required)
    )

    return required, preferred


# ============================================================
# KEYWORDS
# ============================================================

def extract_keywords(jd_text):

    cleaned = clean_for_keywords(jd_text)

    words = re.findall(
        r"[a-zA-Z][a-zA-Z0-9+#.\-]{2,}",
        cleaned
    )

    counts = {}

    for word in words:

        word = word.lower().strip(".-")

        if word in STOP_WORDS:
            continue

        if word in GENERIC_WORDS:
            continue

        if len(word) < 3:
            continue

        counts[word] = counts.get(
            word,
            0
        ) + 1

    ordered = sorted(
        counts.items(),
        key=lambda x: (-x[1], x[0])
    )

    return [
        word
        for word, _ in ordered[:30]
    ]


def keyword_analysis(resume_text, jd_text):

    keywords = extract_keywords(jd_text)

    matching = []
    missing = []

    for keyword in keywords:

        if phrase_present(
            resume_text,
            keyword
        ):
            matching.append(keyword)
        else:
            missing.append(keyword)

    score = (
        round(
            len(matching) /
            len(keywords) * 100
        )
        if keywords
        else 0
    )

    return {
        "keywords": keywords,
        "matching": matching,
        "missing": missing,
        "score": score
    }


# ============================================================
# SKILL MATCH
# ============================================================

def skill_matching(resume_text, jd_text):

    resume_skills = extract_skills(resume_text)
    jd_skills = extract_skills(jd_text)

    required, preferred = (
        detect_required_preferred_skills(
            jd_text,
            jd_skills
        )
    )

    matching_required = [
        skill
        for skill in required
        if skill in resume_skills
    ]

    missing_required = [
        skill
        for skill in required
        if skill not in resume_skills
    ]

    matching_preferred = [
        skill
        for skill in preferred
        if skill in resume_skills
    ]

    missing_preferred = [
        skill
        for skill in preferred
        if skill not in resume_skills
    ]

    required_score = (
        round(
            len(matching_required) /
            len(required) * 100
        )
        if required
        else 100
    )

    preferred_score = (
        round(
            len(matching_preferred) /
            len(preferred) * 100
        )
        if preferred
        else 100
    )

    if required:
        score = (
            round(
                required_score * 0.75
                + preferred_score * 0.25
            )
            if preferred
            else required_score
        )
    else:
        score = preferred_score

    return {
        "resume_skills": resume_skills,
        "jd_skills": jd_skills,
        "required": required,
        "preferred": preferred,
        "matching_required": matching_required,
        "missing_required": missing_required,
        "matching_preferred": matching_preferred,
        "missing_preferred": missing_preferred,
        "required_score": required_score,
        "preferred_score": preferred_score,
        "score": score,
        "categories": categorize_skills(
            resume_skills
        )
    }


# ============================================================
# STRUCTURE
# ============================================================

def analyze_structure(text):

    lower = text.lower()

    sections = {
        "Contact Information": bool(
            re.search(
                r"@|phone|mobile|\+91|\b\d{10}\b",
                lower
            )
        ),

        "Summary / Objective": bool(
            re.search(
                r"summary|objective|profile",
                lower
            )
        ),

        "Skills": bool(
            re.search(
                r"skills|technical skills|core skills",
                lower
            )
        ),

        "Education": bool(
            re.search(
                r"education|academic|qualification",
                lower
            )
        ),

        "Experience": bool(
            re.search(
                r"experience|employment|work history|internship",
                lower
            )
        ),

        "Projects": bool(
            re.search(
                r"projects|project experience",
                lower
            )
        ),

        "Certifications": bool(
            re.search(
                r"certifications|certificates|courses",
                lower
            )
        ),

        "Achievements": bool(
            re.search(
                r"achievements|awards|honors",
                lower
            )
        )
    }

    present = sum(sections.values())

    score = round(
        present / len(sections) * 100
    )

    return {
        "sections": sections,
        "score": score
    }


# ============================================================
# CONTENT
# ============================================================

def analyze_content(text):

    word_count = len(
        re.findall(
            r"\b\w+\b",
            text
        )
    )

    verbs = [
        "developed", "created", "built",
        "designed", "implemented", "optimized",
        "analyzed", "managed", "led",
        "improved", "automated", "integrated",
        "tested", "deployed", "engineered",
        "delivered"
    ]

    action_count = sum(
        len(
            re.findall(
                rf"\b{verb}\b",
                text.lower()
            )
        )
        for verb in verbs
    )

    metrics = re.findall(
        r"\b\d+(?:\.\d+)?\s*(?:%|percent|users|records|projects|days|months|years|hours|x)\b",
        text.lower()
    )

    score = 50

    if word_count >= 250:
        score += 15

    if word_count >= 400:
        score += 10

    if action_count >= 3:
        score += 15

    if metrics:
        score += 10

    return {
        "word_count": word_count,
        "action_verb_count": action_count,
        "metrics": metrics,
        "score": min(score, 100)
    }


# ============================================================
# SCORES
# ============================================================

def calculate_overall_score(
    skill,
    keyword,
    structure,
    content
):
    return round(
        skill * 0.40
        + keyword * 0.25
        + structure * 0.20
        + content * 0.15
    )


def calculate_job_match(
    skill_data,
    keyword_data,
    structure
):

    skill = skill_data.get(
        "score",
        0
    )

    keyword = keyword_data.get(
        "score",
        0
    )

    score = round(
        skill * 0.60
        + keyword * 0.25
        + structure * 0.15
    )

    if score >= 85:
        level = "Excellent Match"
        explanation = (
            "Your resume strongly aligns with the job."
        )

    elif score >= 70:
        level = "Strong Match"
        explanation = (
            "Your resume matches many important requirements."
        )

    elif score >= 55:
        level = "Moderate Match"
        explanation = (
            "Your resume has a good foundation but several areas need improvement."
        )

    else:
        level = "Low Match"
        explanation = (
            "Your resume needs significant tailoring for this role."
        )

    return {
        "score": score,
        "level": level,
        "explanation": explanation
    }


# ============================================================
# NAME / CONTACT / ROLE
# ============================================================

def extract_name(text):

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    for line in lines[:8]:

        clean = re.sub(
            r"[^A-Za-z .'-]",
            "",
            line
        ).strip()

        if (
            2 <= len(clean.split()) <= 5
            and "resume" not in clean.lower()
            and "curriculum" not in clean.lower()
        ):
            return clean

    return "Candidate"


def extract_email(text):

    match = re.search(
        r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
        text
    )

    return match.group(0) if match else ""


def extract_phone(text):

    match = re.search(
        r"(?:\+91[\s-]?)?[6-9]\d{9}",
        text
    )

    return match.group(0) if match else ""


ROLE_WORDS = [
    "developer", "engineer", "analyst", "intern",
    "designer", "scientist", "manager", "programmer",
    "architect", "consultant", "tester"
]


def _clean_role(text):
    text = re.sub(r"[#*_`>]", " ", text)
    text = re.sub(r"\s+", " ", text).strip(" -:|")

    # Cut sentence fragments such as
    # "Junior Python Developer To Help Build".
    text = re.split(
        r"\s+(?:to|who|that|with|for|in|at|and|on)\s+",
        text,
        maxsplit=1,
        flags=re.IGNORECASE
    )[0]

    return text.strip(" -:|,.")


def detect_target_role(jd):

    # 1. Explicit label: "Job title: ...", "Role: ...", "Position: ..."
    match = re.search(
        r"(?:job title|role|position)\s*\**\s*[:\-]\s*\**\s*([^\n]{3,80})",
        jd,
        flags=re.IGNORECASE
    )

    if match:
        role = _clean_role(match.group(1))

        if len(role) > 2:
            return role.title()

    # 2. A heading or first lines that name the role.
    for line in jd.splitlines()[:8]:

        cleaned = _clean_role(line)

        if (
            3 <= len(cleaned) <= 60
            and len(cleaned.split()) <= 6
            and any(
                word in cleaned.lower()
                for word in ROLE_WORDS
            )
        ):
            return cleaned.title()

    # 3. "We are looking for a ... Developer"
    match = re.search(
        r"(?:hiring|looking for|seeking)\s+(?:a|an)?\s*([^\n.,]{3,80})",
        jd,
        flags=re.IGNORECASE
    )

    if match:
        role = _clean_role(match.group(1))

        # Drop leading adjectives like "motivated".
        role = re.sub(
            r"^(?:motivated|passionate|talented|enthusiastic|dedicated)\s+",
            "",
            role,
            flags=re.IGNORECASE
        )

        if (
            len(role) > 2
            and any(
                word in role.lower()
                for word in ROLE_WORDS
            )
        ):
            return role.title()

    return "Target Role"


# ============================================================
# STRENGTHS
# ============================================================

def generate_strengths(data):

    strengths = []

    skill = data.get(
        "skill_score",
        0
    )

    keyword = data.get(
        "keyword_score",
        0
    )

    structure = data.get(
        "structure_score",
        0
    )

    content = data.get(
        "content_score",
        0
    )

    matching = data.get(
        "skill_data",
        {}
    ).get(
        "matching_required",
        []
    )

    if skill >= 75:
        strengths.append(
            "Strong technical skill alignment with the target job."
        )

    elif skill >= 60:
        strengths.append(
            "Good foundation of relevant technical skills."
        )

    if keyword >= 70:
        strengths.append(
            "Good use of job-relevant keywords."
        )

    if structure >= 80:
        strengths.append(
            "Strong and recognizable resume structure."
        )

    if content >= 80:
        strengths.append(
            "Resume content is detailed and readable."
        )

    if len(matching) >= 5:
        strengths.append(
            f"You already match {len(matching)} important required skills."
        )

    if not strengths:
        strengths.append(
            "Your resume provides a foundation that can be improved through targeted tailoring."
        )

    return strengths


# ============================================================
# WEAKNESSES
# ============================================================

def generate_weaknesses(data):

    weaknesses = []

    skill_data = data.get(
        "skill_data",
        {}
    )

    missing_required = skill_data.get(
        "missing_required",
        []
    )

    missing_preferred = skill_data.get(
        "missing_preferred",
        []
    )

    if missing_required:
        weaknesses.append(
            "Missing required skills: "
            + ", ".join(missing_required[:8])
            + "."
        )

    if missing_preferred:
        weaknesses.append(
            "Some preferred skills are missing: "
            + ", ".join(missing_preferred[:6])
            + "."
        )

    if data.get(
        "keyword_score",
        0
    ) < 60:
        weaknesses.append(
            "More job-specific keywords could be included naturally."
        )

    if data.get(
        "structure_score",
        0
    ) < 75:
        weaknesses.append(
            "Some standard resume sections could be improved."
        )

    if data.get(
        "content_score",
        0
    ) < 75:
        weaknesses.append(
            "Use stronger action verbs and measurable results."
        )

    if not weaknesses:
        weaknesses.append(
            "No major weakness detected."
        )

    return weaknesses


# ============================================================
# LEARNING ROADMAP
# ============================================================

def generate_learning_roadmap(data):

    missing = data.get(
        "skill_data",
        {}
    ).get(
        "missing_required",
        []
    )

    roadmap = []

    descriptions = {
        "python":
            "Practice Python syntax, functions, OOP and problem solving.",

        "sql":
            "Practice SELECT, JOIN, GROUP BY, subqueries and database design.",

        "mysql":
            "Practice databases, joins, indexes and queries.",

        "data structures":
            "Study arrays, linked lists, stacks, queues, trees, graphs and hash tables.",

        "algorithms":
            "Practice searching, sorting, recursion, BFS, DFS and Big-O.",

        "oop":
            "Study encapsulation, inheritance, polymorphism and abstraction.",

        "git":
            "Practice commits, branches, merging and conflict resolution.",

        "github":
            "Practice repositories, branches and pull requests.",

        "rest api":
            "Learn HTTP methods, status codes, JSON and API design.",

        "django":
            "Learn models, views, URLs and Django ORM.",

        "flask":
            "Learn routes, templates and REST API development.",

        "react":
            "Learn components, props, state, hooks and API integration.",

        "javascript":
            "Practice ES6, DOM, promises, async/await and APIs.",

        "machine learning":
            "Study preprocessing, supervised learning and model evaluation.",

        "ml":
            "Study classification, regression and model evaluation.",

        "ai":
            "Learn AI fundamentals, ML and generative AI concepts.",

        "docker":
            "Learn images, containers, Dockerfiles and deployment.",

        "aws":
            "Learn EC2, S3, IAM and cloud deployment.",

        "pandas":
            "Practice DataFrames, filtering, grouping and missing values.",

        "numpy":
            "Practice arrays, vectorization and numerical operations."
    }

    for skill in missing[:10]:

        roadmap.append({
            "skill": skill,
            "level": "Priority",
            "action": descriptions.get(
                skill,
                f"Study {skill} and build a small practical project."
            )
        })

    return roadmap


# ============================================================
# TECHNICAL QUESTIONS
# ============================================================

def generate_technical_questions(data):

    skill_data = data.get(
        "skill_data",
        {}
    )

    missing = skill_data.get(
        "missing_required",
        []
    )

    matching = skill_data.get(
        "matching_required",
        []
    )

    questions = []

    for skill in missing:

        skill_questions = TECHNICAL_QUESTIONS.get(
            skill,
            [
                f"What is {skill} and why is it used?",
                f"What are the important concepts of {skill}?",
                f"How would you use {skill} in a project?"
            ]
        )

        for q in skill_questions[:3]:

            questions.append({
                "skill": skill,
                "question": q,
                "priority": "High"
            })

    for skill in matching[:5]:

        skill_questions = TECHNICAL_QUESTIONS.get(
            skill,
            [
                f"How have you used {skill} in your projects?",
                f"What important concepts do you know in {skill}?"
            ]
        )

        for q in skill_questions[:2]:

            questions.append({
                "skill": skill,
                "question": q,
                "priority": "Medium"
            })

    unique = []
    seen = set()

    for item in questions:

        if item["question"] not in seen:
            unique.append(item)
            seen.add(item["question"])

    return unique[:20]


# ============================================================
# HR QUESTIONS
# ============================================================

def generate_hr_questions(data):

    role = data.get(
        "target_role",
        "this role"
    )

    return [
        q.replace(
            "this role",
            role
        )
        for q in HR_QUESTIONS
    ]


# ============================================================
# INTERVIEW PLAN
# ============================================================

def generate_interview_plan(data):

    missing = data.get(
        "skill_data",
        {}
    ).get(
        "missing_required",
        []
    )

    return [
        {
            "day": "Day 1",
            "focus": "Resume & Introduction",
            "tasks": [
                "Practice your 60-second self introduction.",
                "Be able to explain every section of your resume.",
                "Prepare explanations for every project."
            ]
        },
        {
            "day": "Day 2",
            "focus": "Core Technical Skills",
            "tasks": [
                "Revise your strongest technical skills.",
                "Practice basic coding questions.",
                "Review important concepts."
            ]
        },
        {
            "day": "Day 3",
            "focus": "Skill Gaps",
            "tasks": [
                "Study: " + (
                    ", ".join(missing[:5])
                    if missing
                    else "advanced concepts"
                ),
                "Learn the fundamentals.",
                "Build or study one practical example."
            ]
        },
        {
            "day": "Day 4",
            "focus": "Projects",
            "tasks": [
                "Prepare the problem, solution, technologies and result.",
                "Prepare one challenge and achievement for each project.",
                "Practice explaining projects in two minutes."
            ]
        },
        {
            "day": "Day 5",
            "focus": "HR Round",
            "tasks": [
                "Practice Tell Me About Yourself.",
                "Prepare Why should we hire you?",
                "Prepare strengths and weakness answers."
            ]
        },
        {
            "day": "Day 6",
            "focus": "Mock Interview",
            "tasks": [
                "Answer technical questions without notes.",
                "Practice explaining your thought process.",
                "Record yourself answering HR questions."
            ]
        },
        {
            "day": "Day 7",
            "focus": "Final Revision",
            "tasks": [
                "Review the job description.",
                "Review your strongest matching skills.",
                "Review missing skills.",
                "Prepare three questions for the interviewer."
            ]
        }
    ]


# ============================================================
# SUGGESTIONS
# ============================================================

def generate_suggestions(data):

    suggestions = []

    skill_data = data.get(
        "skill_data",
        {}
    )

    missing = skill_data.get(
        "missing_required",
        []
    )

    if missing:
        suggestions.append(
            "Prioritize these missing required skills: "
            + ", ".join(missing[:6])
            + "."
        )

    if data.get(
        "keyword_score",
        0
    ) < 70:
        suggestions.append(
            "Tailor your resume wording to naturally match important job-description keywords."
        )

    if data.get(
        "structure_score",
        0
    ) < 80:
        suggestions.append(
            "Improve resume section structure and headings."
        )

    if data.get(
        "content_score",
        0
    ) < 80:
        suggestions.append(
            "Use stronger action verbs and measurable achievements."
        )

    suggestions.append(
        "Do not add skills you cannot confidently explain in an interview."
    )

    return suggestions


# ============================================================
# PROJECTS
# ============================================================

def generate_project_suggestions(data):

    missing = data.get(
        "skill_data",
        {}
    ).get(
        "missing_required",
        []
    )

    projects = []

    if any(
        x in missing
        for x in ["django", "flask", "fastapi"]
    ):
        projects.append({
            "title": "Python REST API Project",
            "description":
                "Build a backend application using Python, REST APIs and SQL.",
            "skills":
                "Python, REST API, SQL, Django/Flask/FastAPI"
        })

    if any(
        x in missing
        for x in ["sql", "mysql", "postgresql"]
    ):
        projects.append({
            "title": "Database Management Application",
            "description":
                "Create an application that stores and analyzes structured data.",
            "skills":
                "Python, SQL, MySQL"
        })

    if any(
        x in missing
        for x in ["ai", "ml", "machine learning"]
    ):
        projects.append({
            "title": "Machine Learning Prediction System",
            "description":
                "Build an ML project with preprocessing, training and evaluation.",
            "skills":
                "Python, Pandas, NumPy, ML"
        })

    if any(
        x in missing
        for x in ["react", "javascript"]
    ):
        projects.append({
            "title": "Interactive Web Application",
            "description":
                "Build a responsive web application with API integration.",
            "skills":
                "HTML, CSS, JavaScript, React"
        })

    if not projects:
        projects.append({
            "title": "End-to-End Software Project",
            "description":
                "Build a complete application demonstrating your strongest skills.",
            "skills":
                ", ".join(
                    data.get(
                        "skill_data",
                        {}
                    ).get(
                        "matching_required",
                        []
                    )[:8]
                )
        })

    return projects


# ============================================================
# SUMMARY
# ============================================================

def generate_summary(data):

    role = data.get(
        "target_role",
        "software role"
    )

    skills = data.get(
        "skill_data",
        {}
    ).get(
        "matching_required",
        []
    )

    skills_text = ", ".join(
        display_skill(x) for x in skills[:6]
    )

    if not skills_text:
        skills_text = (
            "software development and problem solving"
        )

    return (
        f"Motivated and detail-oriented candidate targeting "
        f"a {role} position, with practical knowledge of "
        f"{skills_text}. Demonstrates strong interest in "
        f"software development, problem solving and building "
        f"practical technology solutions."
    )


# ============================================================
# BULLET IMPROVER
# ============================================================

def improve_bullet(bullet):

    bullet = bullet.strip()

    replacements = {
        "worked on": "Developed",
        "made": "Created",
        "did": "Executed",
        "helped": "Supported",
        "used": "Utilized",
        "responsible for": "Managed",
        "worked with": "Collaborated with"
    }

    improved = bullet

    for old, new in replacements.items():

        if improved.lower().startswith(old):

            improved = (
                new
                + improved[len(old):]
            )

            break

    if improved and not improved.endswith("."):
        improved += "."

    return improved


# ============================================================
# OPTIMIZED RESUME
# ============================================================

SECTION_HEADINGS = {
    "summary": ["professional summary", "summary", "career objective",
                "objective", "profile"],
    "skills": ["technical skills", "skills", "core skills"],
    "projects": ["projects", "relevant projects", "project experience"],
    "experience": ["experience", "work experience", "employment",
                   "internship", "internships", "work history"],
    "education": ["education", "academic background", "qualification",
                  "academics"],
    "certifications": ["certifications", "certificates",
                       "courses & certifications",
                       "courses and certifications", "courses"],
    "achievements": ["achievements", "awards", "honors"],
    "strengths": ["strengths"],
    "languages": ["languages"],
    "interests": ["interests", "hobbies"],
}

ALL_HEADINGS = {
    h for names in SECTION_HEADINGS.values() for h in names
}


def _is_heading(line):
    plain = re.sub(r"[^a-z& ]", "", line.lower()).strip()
    return plain in ALL_HEADINGS


def extract_section(text, section):
    """Return the lines under a heading of the uploaded resume."""

    wanted = SECTION_HEADINGS.get(section, [])
    lines = [line.strip() for line in text.splitlines()]

    start = None

    for i, line in enumerate(lines):
        plain = re.sub(r"[^a-z& ]", "", line.lower()).strip()

        if plain in wanted:
            start = i + 1
            break

    if start is None:
        return []

    collected = []

    for line in lines[start:]:

        if not line:
            continue

        if _is_heading(line):
            break

        collected.append(line)

    return collected


def _as_bullets(lines):
    result = []

    for line in lines:
        is_bullet = bool(re.match(r"^[\u2022\-*\u00b7]\s*", line))
        line = re.sub(r"^[\u2022\-*\u00b7]\s*", "", line).strip()

        if not line:
            continue

        # Only rewrite real bullet sentences, not project titles.
        if is_bullet and len(line.split()) > 5:
            line = improve_bullet(line)

        result.append("\u2022 " + line)

    return result


def generate_optimized_resume(data):

    text = data.get("resume_text", "")

    name = extract_name(text)
    email = extract_email(text)
    phone = extract_phone(text)

    skill_data = data.get("skill_data", {})

    summary = generate_summary(data)

    # Skills that match the job come first, then the rest of the
    # resume's own skills.
    matched = (
        skill_data.get("matching_required", [])
        + skill_data.get("matching_preferred", [])
    )

    skills = list(
        dict.fromkeys(
            matched + skill_data.get("resume_skills", [])
        )
    )

    missing = skill_data.get("missing_required", [])

    contact = " | ".join(
        x for x in [email, phone] if x
    )

    parts = [name]

    if contact:
        parts.append(contact)

    parts.append("")
    parts.append("PROFESSIONAL SUMMARY")
    parts.append(summary)

    parts.append("")
    parts.append("TECHNICAL SKILLS")
    parts.append(
        ", ".join(display_skill(x) for x in skills)
        if skills
        else "Add your technical skills here"
    )

    # Sections are copied from the uploaded resume. A section that
    # the resume does not have is listed at the end instead of being
    # filled with placeholder text.
    missing_sections = []

    for title, key in [
        ("PROJECTS", "projects"),
        ("EXPERIENCE", "experience"),
        ("EDUCATION", "education"),
        ("CERTIFICATIONS", "certifications"),
        ("ACHIEVEMENTS", "achievements"),
    ]:
        lines = extract_section(text, key)

        if not lines:
            if key != "experience":
                missing_sections.append(title.title())
            continue

        parts.append("")
        parts.append(title)

        if key in ("education",):
            parts.extend(lines)
        else:
            parts.extend(_as_bullets(lines))

    languages = extract_section(text, "languages")

    if languages:
        parts.append("")
        parts.append("LANGUAGES")
        parts.append(", ".join(languages))

    notes = []

    if missing_sections:
        notes.append(
            "Sections not found in your resume: "
            + ", ".join(missing_sections)
        )

    if missing:
        notes.append(
            "Skills in the job description that your resume "
            "does not mention: "
            + ", ".join(display_skill(x) for x in missing[:8])
            + ". Add them only if you really have them."
        )

    if notes:
        parts.append("")
        parts.append("NOTES (delete before sending)")
        parts.extend(notes)

    return "\n".join(parts)


# ============================================================
# PDF REPORT
# ============================================================

def create_pdf_report(data):

    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()

    title = ParagraphStyle(
        "CustomTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=22,
        spaceAfter=15
    )

    heading = ParagraphStyle(
        "CustomHeading",
        parent=styles["Heading2"],
        fontSize=14,
        spaceBefore=12,
        spaceAfter=8
    )

    normal = ParagraphStyle(
        "CustomNormal",
        parent=styles["BodyText"],
        fontSize=9,
        leading=13
    )

    story = []

    story.append(
        Paragraph(
            "ResumeAI - Professional ATS Report",
            title
        )
    )

    story.append(
        Paragraph(
            f"<b>Resume:</b> {data.get('filename', 'Resume')}",
            normal
        )
    )

    story.append(
        Paragraph(
            f"<b>Target Role:</b> {data.get('target_role', 'Target Role')}",
            normal
        )
    )

    story.append(Spacer(1, 15))

    story.append(
        Paragraph(
            f"<b>Overall ATS Score: {data.get('overall_score', 0)}%</b>",
            heading
        )
    )

    table_data = [
        ["Category", "Score"],
        ["Skill Match", f"{data.get('skill_score', 0)}%"],
        ["Keyword Match", f"{data.get('keyword_score', 0)}%"],
        ["Structure", f"{data.get('structure_score', 0)}%"],
        ["Content", f"{data.get('content_score', 0)}%"],
        ["Job Match", f"{data.get('job_match', {}).get('score', 0)}%"]
    ]

    table = Table(
        table_data,
        colWidths=[280, 100]
    )

    table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#1f2937")
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),
            (
                "PADDING",
                (0, 0),
                (-1, -1),
                7
            )
        ])
    )

    story.append(table)

    story.append(
        Paragraph(
            "Strengths",
            heading
        )
    )

    for item in data.get(
        "strengths",
        []
    ):
        story.append(
            Paragraph(
                f"• {item}",
                normal
            )
        )

    story.append(
        Paragraph(
            "Areas to Improve",
            heading
        )
    )

    for item in data.get(
        "weaknesses",
        []
    ):
        story.append(
            Paragraph(
                f"• {item}",
                normal
            )
        )

    story.append(
        Paragraph(
            "Missing Required Skills",
            heading
        )
    )

    missing = data.get(
        "skill_data",
        {}
    ).get(
        "missing_required",
        []
    )

    story.append(
        Paragraph(
            ", ".join(missing)
            if missing
            else "None identified",
            normal
        )
    )

    story.append(
        Paragraph(
            "Technical Interview Questions",
            heading
        )
    )

    for i, item in enumerate(
        data.get(
            "technical_questions",
            []
        ),
        1
    ):
        story.append(
            Paragraph(
                f"{i}. {item['question']}",
                normal
            )
        )

    story.append(
        Paragraph(
            "HR Interview Questions",
            heading
        )
    )

    for i, question in enumerate(
        data.get(
            "hr_questions",
            []
        ),
        1
    ):
        story.append(
            Paragraph(
                f"{i}. {question}",
                normal
            )
        )

    story.append(
        Paragraph(
            "Disclaimer",
            heading
        )
    )

    story.append(
        Paragraph(
            "This report is an automated preparation aid. "
            "Actual ATS systems can differ between employers.",
            normal
        )
    )

    doc.build(story)

    buffer.seek(0)

    return buffer.getvalue()


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:

    st.markdown("## 📂 Resume Analysis")

    uploaded_file = st.file_uploader(
        "Upload your resume",
        type=["pdf", "docx", "txt"]
    )

    st.markdown("---")

    st.markdown("### 🎯 Job Description")

    job_description = st.text_area(
        "Paste the job description here",
        height=280,
        placeholder=(
            "Example:\n"
            "Python Developer\n"
            "Required: Python, SQL, Git, REST API..."
        )
    )

    st.markdown("---")

    analyze_button = st.button(
        "🚀 Analyze Resume",
        type="primary",
        use_container_width=True
    )

    if st.session_state.get("analysis"):

        if st.button(
            "🔄 Reset Analysis",
            use_container_width=True
        ):
            st.session_state["analysis"] = None
            st.rerun()

    st.markdown("---")

    st.caption(
        f"ResumeAI {APP_VERSION}"
    )


# ============================================================
# ANALYZE
# ============================================================

if analyze_button:

    if uploaded_file is None:

        st.error(
            "Please upload your resume first."
        )

    elif not job_description.strip():

        st.error(
            "Please paste the job description."
        )

    else:

        with st.spinner(
            "Analyzing your resume..."
        ):

            resume_text = extract_text(
                uploaded_file
            )

            if not resume_text.strip():

                st.error(
                    "Could not extract readable text from the resume."
                )
                st.stop()

            skill_data = skill_matching(
                resume_text,
                job_description
            )

            keyword_data = keyword_analysis(
                resume_text,
                job_description
            )

            structure = analyze_structure(
                resume_text
            )

            content = analyze_content(
                resume_text
            )

            skill_score = skill_data["score"]
            keyword_score = keyword_data["score"]
            structure_score = structure["score"]
            content_score = content["score"]

            overall_score = calculate_overall_score(
                skill_score,
                keyword_score,
                structure_score,
                content_score
            )

            job_match = calculate_job_match(
                skill_data,
                keyword_data,
                structure_score
            )

            target_role = detect_target_role(
                job_description
            )

            base = {
                "filename":
                    uploaded_file.name,

                "resume_text":
                    resume_text,

                "jd_text":
                    job_description,

                "skill_data":
                    skill_data,

                "keyword_data":
                    keyword_data,

                "structure":
                    structure,

                "content":
                    content,

                "skill_score":
                    skill_score,

                "keyword_score":
                    keyword_score,

                "structure_score":
                    structure_score,

                "content_score":
                    content_score,

                "overall_score":
                    overall_score,

                "job_match":
                    job_match,

                "target_role":
                    target_role
            }

            base["strengths"] = (
                generate_strengths(base)
            )

            base["weaknesses"] = (
                generate_weaknesses(base)
            )

            base["suggestions"] = (
                generate_suggestions(base)
            )

            base["project_suggestions"] = (
                generate_project_suggestions(base)
            )

            base["learning_roadmap"] = (
                generate_learning_roadmap(base)
            )

            base["technical_questions"] = (
                generate_technical_questions(base)
            )

            base["hr_questions"] = (
                generate_hr_questions(base)
            )

            base["interview_plan"] = (
                generate_interview_plan(base)
            )

            base["optimized_resume"] = (
                generate_optimized_resume(base)
            )

            base["pdf_bytes"] = (
                create_pdf_report(base)
            )

            st.session_state["analysis"] = {
                "version":
                    APP_VERSION,

                "filename":
                    uploaded_file.name,

                "resume_text":
                    resume_text,

                "jd_text":
                    job_description,

                "skill_data":
                    skill_data,

                "keyword_data":
                    keyword_data,

                "structure":
                    structure,

                "content":
                    content,

                "skill_score":
                    skill_score,

                "keyword_score":
                    keyword_score,

                "structure_score":
                    structure_score,

                "content_score":
                    content_score,

                "overall_score":
                    overall_score,

                "job_match":
                    job_match,

                "target_role":
                    target_role,

                "strengths":
                    base["strengths"],

                "weaknesses":
                    base["weaknesses"],

                "suggestions":
                    base["suggestions"],

                "project_suggestions":
                    base["project_suggestions"],

                "learning_roadmap":
                    base["learning_roadmap"],

                "technical_questions":
                    base["technical_questions"],

                "hr_questions":
                    base["hr_questions"],

                "interview_plan":
                    base["interview_plan"],

                "optimized_resume":
                    base["optimized_resume"],

                "pdf_bytes":
                    base["pdf_bytes"]
            }

        st.rerun()


# ============================================================
# LOAD DATA
# ============================================================

data = st.session_state.get(
    "analysis"
)


if not data:

    st.markdown(
        """
        <div class="info-box">
        <h3>👋 Welcome to ResumeAI</h3>
        <p>
        Upload your resume and paste a job description.
        ResumeAI will analyze ATS compatibility, job matching,
        skill gaps and interview preparation.
        </p>
        </div>
        """,
        unsafe_allow_html=True
    )

    c1, c2, c3 = st.columns(3)

    with c1:
        st.markdown(
            """
            ### 📊 ATS Analysis
            Analyze skills, keywords, structure and content.
            """
        )

    with c2:
        st.markdown(
            """
            ### 🎯 Job Match
            Compare your resume against the target role.
            """
        )

    with c3:
        st.markdown(
            """
            ### 💼 Interview Prep
            Prepare technical and HR questions.
            """
        )

    st.stop()


# ============================================================
# SAFE VALUES
# ============================================================

overall_score = data.get(
    "overall_score",
    0
)

skill_score = data.get(
    "skill_score",
    0
)

keyword_score = data.get(
    "keyword_score",
    0
)

structure_score = data.get(
    "structure_score",
    0
)

content_score = data.get(
    "content_score",
    0
)

skill_data = data.get(
    "skill_data",
    {}
)

keyword_data = data.get(
    "keyword_data",
    {}
)

structure = data.get(
    "structure",
    {}
)

content = data.get(
    "content",
    {}
)

job_match = data.get(
    "job_match",
    {}
)

target_role = data.get(
    "target_role",
    "Target Role"
)


# ============================================================
# STATUS
# ============================================================

if overall_score >= 85:

    message = "Excellent ATS readiness 🎉"

elif overall_score >= 70:

    message = "Strong resume with room for improvement 👍"

elif overall_score >= 55:

    message = "Good foundation — tailor it further 🚀"

else:

    message = "Needs improvement before applying ⚠️"


st.markdown(
    f"""
    <div class="info-box">
    <b>Target Role:</b> {target_role}<br>
    <b>ATS Status:</b> {message}
    </div>
    """,
    unsafe_allow_html=True
)


# ============================================================
# TABS
# ============================================================

tabs = st.tabs([
    "📊 Overview",
    "🎯 Job Match",
    "🧠 Skills",
    "✍️ AI Improvement",
    "🚀 Projects",
    "📄 Optimized Resume",
    "💼 Interview Prep",
    "📈 Analytics",
    "📋 Structure",
    "📄 Resume Text"
])


# ============================================================
# OVERVIEW
# ============================================================

with tabs[0]:

    st.markdown(
        '<div class="section-title">📊 ATS Overview</div>',
        unsafe_allow_html=True
    )

    c1, c2, c3, c4, c5 = st.columns(5)

    with c1:
        st.metric(
            "ATS Score",
            f"{overall_score}%"
        )

    with c2:
        st.metric(
            "Skill Match",
            f"{skill_score}%"
        )

    with c3:
        st.metric(
            "Keywords",
            f"{keyword_score}%"
        )

    with c4:
        st.metric(
            "Structure",
            f"{structure_score}%"
        )

    with c5:
        st.metric(
            "Job Match",
            f"{job_match.get('score', 0)}%"
        )

    st.markdown("### 📈 ATS Score")

    st.progress(
        max(0, min(100, overall_score)) / 100
    )

    st.markdown("### 💪 Resume Strengths")

    for item in data.get(
        "strengths",
        []
    ):

        st.markdown(
            f"""
            <div class="success-box">
            ✅ {item}
            </div>
            """,
            unsafe_allow_html=True
        )

    st.markdown("### ⚠️ Areas to Improve")

    for item in data.get(
        "weaknesses",
        []
    ):

        st.markdown(
            f"""
            <div class="warning-box">
            ⚠️ {item}
            </div>
            """,
            unsafe_allow_html=True
        )

    st.download_button(
        "⬇️ Download Professional PDF Report",
        data=data.get(
            "pdf_bytes",
            b""
        ),
        file_name="ResumeAI_ATS_Analysis_Report.pdf",
        mime="application/pdf",
        use_container_width=True
    )


# ============================================================
# JOB MATCH
# ============================================================

with tabs[1]:

    st.markdown(
        '<div class="section-title">🎯 Job Match Analysis</div>',
        unsafe_allow_html=True
    )

    match_score = job_match.get(
        "score",
        0
    )

    st.metric(
        "Job Match Score",
        f"{match_score}%",
        job_match.get(
            "level",
            ""
        )
    )

    st.progress(
        max(0, min(100, match_score)) / 100
    )

    st.info(
        job_match.get(
            "explanation",
            ""
        )
    )

    st.markdown(
        "### ✅ Matching Required Skills"
    )

    matching_required = skill_data.get(
        "matching_required",
        []
    )

    if matching_required:

        st.write(
            " • ".join(
                display_skill(x)
                for x in matching_required
            )
        )

    else:

        st.warning(
            "No matching required skills detected."
        )

    st.markdown(
        "### ❌ Missing Required Skills"
    )

    missing_required = skill_data.get(
        "missing_required",
        []
    )

    if missing_required:

        for skill in missing_required:

            st.markdown(
                f"""
                <div class="danger-box">
                ❌ <b>{display_skill(skill)}</b>
                </div>
                """,
                unsafe_allow_html=True
            )

    else:

        st.success(
            "No major required skill gaps detected."
        )

    c1, c2 = st.columns(2)

    with c1:

        st.markdown(
            "### ⭐ Matching Preferred"
        )

        for skill in skill_data.get(
            "matching_preferred",
            []
        ):
            st.success(
                display_skill(skill)
            )

    with c2:

        st.markdown(
            "### ⭐ Missing Preferred"
        )

        for skill in skill_data.get(
            "missing_preferred",
            []
        ):
            st.warning(
                display_skill(skill)
            )


# ============================================================
# SKILLS
# ============================================================

with tabs[2]:

    st.markdown(
        '<div class="section-title">🧠 Skills Intelligence</div>',
        unsafe_allow_html=True
    )

    categories = skill_data.get(
        "categories",
        {}
    )

    for category, skills in categories.items():

        st.markdown(
            f"### {category}"
        )

        for skill in skills:

            st.markdown(
                f"""
                <div class="skill-box">
                💡 {display_skill(skill)}
                </div>
                """,
                unsafe_allow_html=True
            )

    st.markdown(
        "### 🎯 Skills to Prioritize"
    )

    roadmap = data.get(
        "learning_roadmap",
        []
    )

    if roadmap:

        for i, item in enumerate(
            roadmap,
            1
        ):

            st.markdown(
                f"""
                <div class="warning-box">
                <b>{i}. {display_skill(item['skill'])}</b>
                <br>
                <span class="small-muted">
                {item['level']} • {item['action']}
                </span>
                </div>
                """,
                unsafe_allow_html=True
            )

    else:

        st.success(
            "No major skill gaps detected."
        )


# ============================================================
# AI IMPROVEMENT
# ============================================================

with tabs[3]:

    st.markdown(
        '<div class="section-title">✍️ AI Resume Improvement</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        "### 💡 Personalized Recommendations"
    )

    for item in data.get(
        "suggestions",
        []
    ):

        st.markdown(
            f"""
            <div class="info-box">
            💡 {item}
            </div>
            """,
            unsafe_allow_html=True
        )

    st.markdown(
        "### 📝 Resume Bullet Point Improver"
    )

    bullet = st.text_area(
        "Paste one resume bullet point",
        placeholder=(
            "Example: Worked on a student management system"
        )
    )

    if st.button(
        "✨ Improve Bullet Point"
    ):

        improved = improve_bullet(
            bullet
        )

        if improved:

            st.success(
                "Improved bullet point:"
            )

            st.write(
                improved
            )

        else:

            st.warning(
                "Please enter a bullet point."
            )


# ============================================================
# PROJECTS
# ============================================================

with tabs[4]:

    st.markdown(
        '<div class="section-title">🚀 Project Recommendations</div>',
        unsafe_allow_html=True
    )

    st.info(
        "Projects are selected based on skills required by the target job."
    )

    for project in data.get(
        "project_suggestions",
        []
    ):

        st.markdown(
            f"""
            <div class="skill-box">
            <h3>🚀 {project['title']}</h3>
            <p>{project['description']}</p>
            <b>Skills:</b> {project['skills']}
            </div>
            """,
            unsafe_allow_html=True
        )


# ============================================================
# OPTIMIZED RESUME
# ============================================================

with tabs[5]:

    st.markdown(
        '<div class="section-title">📄 Optimized Resume</div>',
        unsafe_allow_html=True
    )

    st.info(
        "Review the generated draft and only keep information that is truthful."
    )

    optimized = data.get(
        "optimized_resume",
        ""
    )

    st.text_area(
        "Optimized Resume Draft",
        value=optimized,
        height=650
    )

    st.download_button(
        "⬇️ Download Optimized Resume",
        data=optimized,
        file_name="Optimized_Resume.txt",
        mime="text/plain",
        use_container_width=True
    )


# ============================================================
# INTERVIEW PREP
# ============================================================

with tabs[6]:

    st.markdown(
        '<div class="section-title">💼 Interview Intelligence</div>',
        unsafe_allow_html=True
    )

    st.info(
        f"Target role: {target_role}"
    )

    interview_tabs = st.tabs([
        "💻 Technical",
        "👩‍💼 HR",
        "📅 7-Day Plan"
    ])

    with interview_tabs[0]:

        st.markdown(
            "### 💻 Technical Interview Questions"
        )

        questions = data.get(
            "technical_questions",
            []
        )

        for i, item in enumerate(
            questions,
            1
        ):

            icon = (
                "🔴"
                if item["priority"] == "High"
                else "🟡"
            )

            st.markdown(
                f"""
                <div class="question-card">
                <b>{i}. {item['question']}</b>
                <br><br>
                {icon}
                <b>{display_skill(item['skill'])}</b>
                —
                {item['priority']} Priority
                </div>
                """,
                unsafe_allow_html=True
            )

        st.info(
            "Be prepared to explain where you used every technical skill listed on your resume."
        )

    with interview_tabs[1]:

        st.markdown(
            "### 👩‍💼 HR Interview Questions"
        )

        for i, question in enumerate(
            data.get(
                "hr_questions",
                []
            ),
            1
        ):

            st.markdown(
                f"""
                <div class="question-card">
                <b>{i}. {question}</b>
                </div>
                """,
                unsafe_allow_html=True
            )

        st.markdown(
            "### ⭐ Tell Me About Yourself Framework"
        )

        st.write(
            """
            1. Who you are academically/professionally  
            2. Your strongest technical skills  
            3. One or two relevant projects  
            4. What role you are targeting  
            5. Why you are interested in the opportunity  

            Keep it around **60–90 seconds**.
            """
        )

    with interview_tabs[2]:

        st.markdown(
            "### 📅 7-Day Interview Preparation Plan"
        )

        for day in data.get(
            "interview_plan",
            []
        ):

            st.markdown(
                f"#### {day['day']} — {day['focus']}"
            )

            for task in day["tasks"]:

                st.checkbox(
                    task,
                    key=f"{day['day']}_{task}"
                )


# ============================================================
# ANALYTICS
# ============================================================

with tabs[7]:

    st.markdown(
        '<div class="section-title">📈 Resume Analytics</div>',
        unsafe_allow_html=True
    )

    # IMPORTANT:
    # NO st.bar_chart() HERE.
    # It requires Pandas, which is causing the DLL error.

    st.markdown(
        "### 📊 Score Breakdown"
    )

    analytics = [
        ("ATS Score", overall_score),
        ("Skill Match", skill_score),
        ("Keyword Match", keyword_score),
        ("Structure", structure_score),
        ("Content", content_score),
        ("Job Match", job_match.get("score", 0))
    ]

    for label, score in analytics:

        st.markdown(
            f"**{label} — {score}%**"
        )

        st.progress(
            max(0, min(100, score)) / 100
        )

    st.markdown(
        "### 📌 Score Interpretation"
    )

    if overall_score >= 85:

        st.success(
            "85–100%: Excellent ATS readiness."
        )

    elif overall_score >= 70:

        st.info(
            "70–84%: Strong resume with some improvements possible."
        )

    elif overall_score >= 55:

        st.warning(
            "55–69%: Good foundation, but improvements are recommended."
        )

    else:

        st.error(
            "Below 55%: Significant tailoring is recommended."
        )

    st.markdown(
        "### 🔍 Content Statistics"
    )

    c1, c2, c3 = st.columns(3)

    with c1:

        st.metric(
            "Resume Words",
            content.get(
                "word_count",
                0
            )
        )

    with c2:

        st.metric(
            "Action Verbs",
            content.get(
                "action_verb_count",
                0
            )
        )

    with c3:

        st.metric(
            "Metrics Found",
            len(
                content.get(
                    "metrics",
                    []
                )
            )
        )

    st.markdown(
        "### 🔑 Keyword Analytics"
    )

    c1, c2 = st.columns(2)

    with c1:

        st.markdown(
            "#### ✅ Matching Keywords"
        )

        for keyword in keyword_data.get(
            "matching",
            []
        )[:15]:

            st.success(
                keyword
            )

    with c2:

        st.markdown(
            "#### ❌ Missing Keywords"
        )

        for keyword in keyword_data.get(
            "missing",
            []
        )[:15]:

            st.warning(
                keyword
            )


# ============================================================
# STRUCTURE
# ============================================================

with tabs[8]:

    st.markdown(
        '<div class="section-title">📋 Resume Structure</div>',
        unsafe_allow_html=True
    )

    st.metric(
        "Structure Score",
        f"{structure_score}%"
    )

    st.progress(
        max(0, min(100, structure_score)) / 100
    )

    sections = structure.get(
        "sections",
        {}
    )

    for section, present in sections.items():

        if present:

            st.success(
                f"✅ {section}"
            )

        else:

            st.warning(
                f"⚠️ {section} — not detected"
            )

    st.markdown(
        "### ✍️ Content Quality"
    )

    st.metric(
        "Content Score",
        f"{content_score}%"
    )

    st.write(
        f"**Word count:** {content.get('word_count', 0)}"
    )

    st.write(
        f"**Action verbs:** {content.get('action_verb_count', 0)}"
    )

    metrics = content.get(
        "metrics",
        []
    )

    if metrics:

        st.write(
            "**Metrics detected:** "
            + ", ".join(metrics)
        )

    else:

        st.warning(
            "Try adding measurable results such as percentages, users, records or time saved."
        )


# ============================================================
# RESUME TEXT
# ============================================================

with tabs[9]:

    st.markdown(
        '<div class="section-title">📄 Extracted Resume Text</div>',
        unsafe_allow_html=True
    )

    resume_text = data.get(
        "resume_text",
        ""
    )

    st.text_area(
        "Resume Text",
        value=resume_text,
        height=600
    )

    st.download_button(
        "⬇️ Download Extracted Resume Text",
        data=resume_text,
        file_name="Extracted_Resume_Text.txt",
        mime="text/plain",
        use_container_width=True
    )


# ============================================================
# FOOTER
# ============================================================

st.markdown("---")

st.markdown(
    """
    <div style="text-align:center; color:#6b7280; padding:20px;">
        <b>ResumeAI V12.1</b><br>
        ATS Analysis • Job Matching • Resume Optimization • Interview Intelligence
        <br><br>
        Built as an AI-powered placement project 🚀
    </div>
    """,
    unsafe_allow_html=True
)
